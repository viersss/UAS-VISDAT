import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

# Add title in single column section
doc.add_heading("Test Title", level=0)

# Add new section for two columns
sec2 = doc.add_section()
sectPr = sec2._sectPr
cols = sectPr.xpath('./w:cols')
if cols:
    cols[0].set(qn('w:num'), '2')
    cols[0].set(qn('w:space'), '360') # 0.25 inch space between columns
else:
    cols_elm = parse_xml(r'<w:cols xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:num="2" w:space="360"/>')
    sectPr.append(cols_elm)

doc.add_paragraph("This is column 1 text. " * 30)

doc.save("test_col.docx")
print("Saved test_col.docx successfully!")
